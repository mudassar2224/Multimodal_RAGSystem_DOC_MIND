import io
import subprocess
import tempfile
from dataclasses import asdict
from pathlib import Path

import imageio_ffmpeg
import pandas as pd
import pymupdf
from docx import Document as DocxDocument
from pptx import Presentation

from docmind.core.schemas import ChunkDraft, SourceRef
from docmind.ingestion.chunking import chunk_text


def parse_pdf(data: bytes, source: SourceRef) -> tuple[list[ChunkDraft], list[tuple[str, bytes]]]:
    pdf = pymupdf.open(stream=data, filetype="pdf")
    chunks: list[ChunkDraft] = []
    images: list[tuple[str, bytes]] = []
    for index, page in enumerate(pdf, start=1):
        page_source = SourceRef(**{**asdict(source), "page": index})
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
        image = pixmap.tobytes("png")
        key = f"derived/{source.document_id}/page-{index}.png"
        images.append((key, image))
        for text in chunk_text(page.get_text()):
            chunks.append(ChunkDraft(text=text, source=page_source, visual_object_key=key))
    return chunks, images


def parse_docx(data: bytes, source: SourceRef) -> list[ChunkDraft]:
    document = DocxDocument(io.BytesIO(data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    return [ChunkDraft(text=part, source=source) for part in chunk_text(text)]


def parse_pptx(data: bytes, source: SourceRef) -> list[ChunkDraft]:
    presentation = Presentation(io.BytesIO(data))
    drafts: list[ChunkDraft] = []
    for number, slide in enumerate(presentation.slides, start=1):
        slide_source = SourceRef(document_id=source.document_id, filename=source.filename, slide=number)
        text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
        drafts.extend(ChunkDraft(text=part, source=slide_source) for part in chunk_text(text))
    return drafts


def parse_tabular(data: bytes, suffix: str, source: SourceRef) -> list[ChunkDraft]:
    if suffix == ".csv":
        frames = {"CSV": pd.read_csv(io.BytesIO(data))}
    elif suffix == ".json":
        frames = {"JSON": pd.read_json(io.BytesIO(data))}
    else:
        frames = pd.read_excel(io.BytesIO(data), sheet_name=None)
    drafts: list[ChunkDraft] = []
    for sheet, frame in frames.items():
        for start in range(0, len(frame), 40):
            block = frame.iloc[start:start + 40]
            end = start + len(block) + 1
            source_ref = SourceRef(document_id=source.document_id, filename=source.filename, sheet=str(sheet), cell_range=f"A{start + 2}:{_column_name(len(frame.columns))}{end}")
            drafts.append(ChunkDraft(text=block.to_csv(index=False), source=source_ref))
    return drafts


def _column_name(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result or "A"


def extract_video_assets(data: bytes, suffix: str, interval: int) -> tuple[bytes | None, list[tuple[float, bytes]]]:
    """Uses server-side ffmpeg only; clients never require local media tooling or a GPU."""
    with tempfile.TemporaryDirectory() as directory:
        src = Path(directory) / f"input{suffix}"
        audio = Path(directory) / "audio.wav"
        src.write_bytes(data)
        ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
        # A video can legitimately have no audio track. Index frames even in that case.
        audio_result = subprocess.run(
            [ffmpeg, "-y", "-i", str(src), "-vn", "-acodec", "pcm_s16le", str(audio)],
            check=False,
            capture_output=True,
        )
        pattern = Path(directory) / "frame-%06d.png"
        subprocess.run([ffmpeg, "-y", "-i", str(src), "-vf", f"fps=1/{interval}", str(pattern)], check=True, capture_output=True)
        frame_paths = sorted(Path(directory).glob("frame-*.png"))
        if not frame_paths:
            # Short clips can contain no exact interval boundary; always retain their opening frame.
            subprocess.run([ffmpeg, "-y", "-i", str(src), "-frames:v", "1", str(pattern)], check=True, capture_output=True)
            frame_paths = sorted(Path(directory).glob("frame-*.png"))
        return (
            audio.read_bytes() if audio_result.returncode == 0 and audio.exists() else None,
            [(index * interval, path.read_bytes()) for index, path in enumerate(frame_paths)],
        )
